"""Convert the Markdown summary into a Word document."""
import io
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def _add_inline(par, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            par.add_run(part[1:-1]).font.name = "Consolas"
        elif part.startswith("[") and "](" in part:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            run = par.add_run(f"{label} ({url})")
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x9A)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def _table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in r)]
    if not cells:
        return
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(cells):
        row = t.add_row().cells
        for j in range(ncol):
            p = row[j].paragraphs[0]
            _add_inline(p, r[j] if j < len(r) else "")
            for run in p.runs:
                run.font.size = Pt(8.5)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def markdown_to_docx(md: str) -> bytes:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            _table(doc, block)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            doc.add_heading(m.group(2).strip(), level=min(len(m.group(1)) - 1, 3) if len(m.group(1)) > 1 else 0)
        elif re.match(r"^[-*]\s+", s):
            _add_inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^[-*]\s+", "", s))
        elif re.match(r"^\d+[.)]\s+", s):
            _add_inline(doc.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s+", "", s))
        elif s.startswith(">"):
            p = doc.add_paragraph()
            _add_inline(p, s.lstrip("> ").strip())
            for run in p.runs:
                run.italic = True
        elif s in ("---", "***"):
            pass
        elif s:
            _add_inline(doc.add_paragraph(), s)
        i += 1
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
